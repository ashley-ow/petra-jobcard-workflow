from __future__ import annotations

import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.workbook import Workbook
from openpyxl.workbook.properties import CalcProperties

from .models import UploadedFileRecord, ValidatedJobPlan
from .storage import safe_filename

_REQUIRED_NAMES = {
    "ProjectName",
    "ProjectDescription",
    "GeneratedAt",
    "SourcePdf",
    "ReviewNotes",
    "JobTableStart",
    "TotalMandays",
}


class ExcelJobCardWriter:
    """Populate a named-range-driven Excel template while preserving its formatting."""

    def __init__(self, template_path: Path, output_dir: Path, task_capacity: int) -> None:
        self.template_path = template_path
        self.output_dir = output_dir
        self.task_capacity = task_capacity

    def write(
        self,
        plan: ValidatedJobPlan,
        request_description: str,
        source: UploadedFileRecord,
    ) -> Path:
        if not self.template_path.exists():
            raise FileNotFoundError(f"Excel template not found: {self.template_path}")
        if len(plan.tasks) > self.task_capacity:
            raise ValueError(
                f"The plan has {len(plan.tasks)} tasks but the template capacity is {self.task_capacity}."
            )

        self.output_dir.mkdir(parents=True, exist_ok=True)
        stem = safe_filename(plan.project_title, fallback="job-card")
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S-%f")
        output = self.output_dir / f"{stem}-{timestamp}.xlsx"
        shutil.copy2(self.template_path, output)

        workbook = load_workbook(output)
        self._validate_names(workbook)
        project_cell = named_cell(workbook, "ProjectName")
        description_cell = named_cell(workbook, "ProjectDescription")
        generated_cell = named_cell(workbook, "GeneratedAt")
        source_cell = named_cell(workbook, "SourcePdf")
        review_cell = named_cell(workbook, "ReviewNotes")
        total_cell = named_cell(workbook, "TotalMandays")
        start_cell = named_cell(workbook, "JobTableStart")
        sheet = start_cell.parent
        start_row, start_column = start_cell.row, start_cell.column

        project_cell.value = plan.project_title
        description_cell.value = request_description
        generated_cell.value = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
        source_cell.value = source.original_name
        notes = []
        if plan.assumptions:
            notes.append("Assumptions: " + " | ".join(plan.assumptions))
        if plan.warnings:
            notes.append("Warnings: " + " | ".join(plan.warnings))
        review_cell.value = "\n".join(notes) if notes else "No workflow warnings."

        for offset in range(self.task_capacity):
            row = start_row + offset
            for column in range(start_column, start_column + 10):
                sheet.cell(row=row, column=column).value = None

        for offset, task in enumerate(plan.tasks):
            row = start_row + offset
            values = [
                offset + 1,
                task.job_task,
                task.matched_reference_task or "UNMATCHED",
                task.quantity,
                task.unit,
                task.unit_mandays,
                None,
                task.source_page,
                f"{task.source_quote or 'No reference quote'}\n{task.rationale}",
                "REVIEW" if task.review_required else "OK",
            ]
            for column_offset, value in enumerate(values):
                cell = sheet.cell(row=row, column=start_column + column_offset)
                cell.value = value
            total = sheet.cell(row=row, column=start_column + 6)
            quantity_ref = sheet.cell(row=row, column=start_column + 3).coordinate
            unit_md_ref = sheet.cell(row=row, column=start_column + 5).coordinate
            total.value = f"=ROUND({quantity_ref}*{unit_md_ref},2)"

        total_column_letter = sheet.cell(row=start_row, column=start_column + 6).column_letter
        total_cell.value = (
            f"=SUM('{sheet.title}'!{total_column_letter}{start_row}:"
            f"{total_column_letter}{start_row + self.task_capacity - 1})"
        )
        if workbook.calculation is None:
            workbook.calculation = CalcProperties()
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        workbook.save(output)
        return output

    def _validate_names(self, workbook: Workbook) -> None:
        existing = set(workbook.defined_names)
        missing = _REQUIRED_NAMES - existing
        if missing:
            raise ValueError(
                "Excel template is missing required named ranges: " + ", ".join(sorted(missing))
            )


def named_cell(workbook: Workbook, name: str) -> Cell:
    definition = workbook.defined_names.get(name)
    if definition is None:
        raise KeyError(f"Missing named range: {name}")
    destinations = list(definition.destinations)
    if len(destinations) != 1:
        raise ValueError(f"Named range '{name}' must refer to exactly one worksheet cell.")
    sheet_name, cell_reference = destinations[0]
    top_left = cell_reference.replace("$", "").split(":", maxsplit=1)[0]
    return workbook[sheet_name][top_left]
