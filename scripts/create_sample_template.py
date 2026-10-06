from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName


DARK = "17324D"
BLUE = "0067B8"
LIGHT_BLUE = "E8F3FA"
PALE = "F4F7F9"
WHITE = "FFFFFF"
GRAY = "66737F"
GREEN = "107C10"
ORANGE = "F5A623"


def build_template(output: Path, capacity: int = 40) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Job Card"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A12"

    widths = {"A": 7, "B": 28, "C": 26, "D": 11, "E": 12, "F": 13, "G": 14, "H": 11, "I": 48, "J": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A1:J2")
    sheet["A1"] = "PROJECT JOB CARD"
    sheet["A1"].font = Font(name="Aptos Display", size=24, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=DARK)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 16

    labels = {
        "A4": "PROJECT",
        "G4": "GENERATED (UTC)",
        "A5": "SOURCE PDF",
        "G5": "TOTAL MANDAYS",
        "A6": "DESCRIPTION",
        "A8": "REVIEW NOTES",
    }
    for coordinate, value in labels.items():
        sheet[coordinate] = value
        sheet[coordinate].font = Font(name="Aptos", size=9, bold=True, color=GRAY)
        sheet[coordinate].alignment = Alignment(vertical="center")

    sheet.merge_cells("B4:F4")
    sheet.merge_cells("H4:J4")
    sheet.merge_cells("B5:F5")
    sheet.merge_cells("H5:J5")
    sheet.merge_cells("B6:J6")
    sheet.merge_cells("B8:J9")
    for coordinate in ["B4", "H4", "B5", "H5", "B6", "B8"]:
        cell = sheet[coordinate]
        cell.fill = PatternFill("solid", fgColor=PALE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.font = Font(name="Aptos", size=10, color=BLUE if coordinate != "H5" else GREEN)
    sheet["H5"].font = Font(name="Aptos Display", size=16, bold=True, color=GREEN)
    sheet.row_dimensions[6].height = 42
    sheet.row_dimensions[8].height = 28
    sheet.row_dimensions[9].height = 28

    headers = [
        "ITEM",
        "PROJECT / JOB TASK",
        "MATCHED REFERENCE TASK",
        "QTY",
        "UNIT",
        "MD / UNIT",
        "TOTAL MD",
        "REF PAGE",
        "EVIDENCE / RATIONALE",
        "STATUS",
    ]
    header_row = 11
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 30

    thin_gray = Side(style="thin", color="CFD8DF")
    for offset in range(capacity):
        row = 12 + offset
        sheet.row_dimensions[row].height = 45
        for column in range(1, 11):
            cell = sheet.cell(row=row, column=column)
            cell.fill = PatternFill("solid", fgColor=WHITE if offset % 2 == 0 else PALE)
            cell.font = Font(name="Aptos", size=9, color=BLUE if column in {2, 4, 5, 6} else "202A33")
            cell.alignment = Alignment(
                horizontal="center" if column in {1, 4, 6, 7, 8, 10} else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin_gray)
        sheet.cell(row=row, column=4).number_format = "0.00"
        sheet.cell(row=row, column=6).number_format = "0.00"
        sheet.cell(row=row, column=7).number_format = "0.00"

    total_row = 12 + capacity + 1
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
    sheet.cell(row=total_row, column=1, value="TOTAL MANDAYS")
    sheet.cell(row=total_row, column=7, value=f"=SUM(G12:G{11 + capacity})")
    for column in range(1, 8):
        cell = sheet.cell(row=total_row, column=column)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="right" if column == 7 else "left", vertical="center")
    sheet.cell(row=total_row, column=7).number_format = "0.00"
    sheet.row_dimensions[total_row].height = 28

    foot = total_row + 2
    sheet.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=10)
    sheet.cell(row=foot, column=1, value="Blue text = workflow input · Formula totals recalculate when opened in Excel · REVIEW rows require human confirmation")
    sheet.cell(row=foot, column=1).font = Font(name="Aptos", size=8, italic=True, color=GRAY)

    names = {
        "ProjectName": "'Job Card'!$B$4",
        "GeneratedAt": "'Job Card'!$H$4",
        "SourcePdf": "'Job Card'!$B$5",
        "TotalMandays": "'Job Card'!$H$5",
        "ProjectDescription": "'Job Card'!$B$6",
        "ReviewNotes": "'Job Card'!$B$8",
        "JobTableStart": "'Job Card'!$A$12",
    }
    for name, target in names.items():
        workbook.defined_names.add(DefinedName(name, attr_text=target))

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


if __name__ == "__main__":
    build_template(Path("templates/job_card_template.xlsx"))
